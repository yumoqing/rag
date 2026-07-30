"""
RAG 文件处理管线 — 文档解析 / 人脸检测 / 声纹处理
GPU 服务映射:
  人脸: https://media.opencomputing.net:10443/face/api/detect
  声纹: https://media.opencomputing.net:10443/voiceprint/extract/submit
"""
import json, os, uuid, base64

# File type classification
TEXT_EXTS = {'.txt', '.md', '.csv', '.json', '.xml', '.html', '.htm', '.py', '.js', '.css', '.yaml', '.yml', '.log', '.rst'}
DOC_EXTS = {'.pdf', '.docx', '.pptx', '.xlsx', '.doc', '.ppt', '.xls'}
IMAGE_EXTS = {'.jpg', '.jpeg', '.png', '.bmp', '.gif', '.webp'}
AUDIO_EXTS = {'.mp3', '.wav', '.flac', '.ogg', '.m4a', '.aac'}
VIDEO_EXTS = {'.mp4', '.avi', '.mov', '.mkv', '.webm'}

# GPU service endpoints
FACE_DETECT_URL = "https://media.opencomputing.net:10443/face/api/detect"
FACE_RECOGNIZE_URL = "https://media.opencomputing.net:10443/face/api/recognize"
VOICEPRINT_EXTRACT_URL = "https://media.opencomputing.net:10443/voiceprint/extract/submit"
VOICEPRINT_STATUS_URL = "https://media.opencomputing.net:10443/voiceprint/extract/status"


def classify_file(file_name):
    ext = os.path.splitext(file_name)[1].lower()
    if ext in TEXT_EXTS:
        return 'text'
    if ext in DOC_EXTS:
        return 'document'
    if ext in IMAGE_EXTS:
        return 'image'
    if ext in AUDIO_EXTS:
        return 'audio'
    if ext in VIDEO_EXTS:
        return 'video'
    return 'other'


async def extract_text(file_data, file_name):
    ext = os.path.splitext(file_name)[1].lower()
    text = ""
    
    if ext in TEXT_EXTS:
        text = file_data.decode('utf-8', errors='replace')
    
    elif ext == '.pdf':
        try:
            import io
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(file_data))
            text = '\n'.join(p.extract_text() or '' for p in reader.pages)
        except Exception as e:
            text = "[PDF error: " + str(e) + "]"
    
    elif ext == '.docx':
        try:
            import io
            from docx import Document
            doc = Document(io.BytesIO(file_data))
            text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            text = "[DOCX error: " + str(e) + "]"
    
    elif ext == '.pptx':
        try:
            import io
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_data))
            slides = []
            for slide in prs.slides:
                parts = [s.text for s in slide.shapes if hasattr(s, 'text') and s.text.strip()]
                slides.append(' '.join(parts))
            text = '\n'.join(slides)
        except Exception as e:
            text = "[PPTX error: " + str(e) + "]"
    
    elif ext in {'.xlsx', '.xls'}:
        try:
            import io
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(file_data), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    r = ' | '.join(str(c) if c is not None else '' for c in row)
                    if r.strip():
                        rows.append(r)
            text = '\n'.join(rows)
        except Exception as e:
            text = "[XLSX error: " + str(e) + "]"
    
    return text.strip()


async def ingest_to_rag(env, text, kb_id, doc_id):
    try:
        from pipeline import ingest as pipeline_ingest
        result = pipeline_ingest(
            text, pipeline_name="kg-rag-standard",
            collection=kb_id, graph_name=kb_id, llm_func=None)
        return result
    except Exception as e:
        return {"error": str(e), "chunks": 0}


async def detect_faces(file_data, file_name):
    """Detect faces via GPU face-service (InsightFace buffalo_l)"""
    import aiohttp
    try:
        b64 = base64.b64encode(file_data).decode()
        payload = {"image": b64}
        async with aiohttp.ClientSession() as session:
            async with session.post(
                FACE_DETECT_URL,
                json=payload,
                timeout=aiohttp.ClientTimeout(total=30)
            ) as resp:
                if resp.status == 200:
                    result = await resp.json()
                    return {
                        "faces": result.get("faces", result.get("count", 0)),
                        "details": result
                    }
                return {"error": "face service returned " + str(resp.status), "faces": 0}
    except Exception as e:
        return {"error": str(e), "faces": 0}


async def extract_voiceprint(file_data, file_name):
    import aiohttp
    try:
        async with aiohttp.ClientSession() as session:
            form = aiohttp.FormData()
            form.add_field('file', file_data, filename=file_name, content_type='audio/wav')
            async with session.post(VOICEPRINT_EXTRACT_URL, data=form, timeout=aiohttp.ClientTimeout(total=60)) as resp:
                if resp.status == 200:
                    r = await resp.json()
                    if r.get('status') == 'SUCCEEDED':
                        return {'voiceprint': r, 'speakers': 1}
                    return {'error': r.get('error', ''), 'speakers': 0}
                return {'error': 'service returned ' + str(resp.status), 'speakers': 0}
    except Exception as e:
        return {'error': str(e), 'speakers': 0}

async def process_upload(env, file_data, kb_id, folder_id, file_name):
    """Full upload pipeline — save + classify + process + DB record"""
    import uuid as _uuid
    doc_id = str(_uuid.uuid4()).hex[:16]
    ext = '.' + file_name.rsplit('.', 1)[1] if '.' in file_name else '.bin'
    saved_name = doc_id + ext
    file_path = '/d/rag/ragserver/pkgs/rag/rag/files/' + saved_name
    with open(file_path, 'wb') as f:
        f.write(file_data)
    file_size = len(file_data)

    ft = classify_file(file_name)
    ingest_result = None
    face_result = None
    voice_result = None
    status = 'pending'

    if ft in ('text', 'document'):
        text = await extract_text(file_data, file_name)
        if text and len(text) > 10:
            ingest_result = await ingest_to_rag(env, text, kb_id, doc_id)
            status = 'done' if ingest_result and not ingest_result.get('error') else 'error'
        else:
            status = 'done'
    elif ft == 'image':
        face_result = await detect_faces(file_data, file_name)
        status = 'done'
    elif ft == 'audio':
        voice_result = await extract_voiceprint(file_data, file_name)
        status = 'done'
    elif ft == 'video':
        face_result = await detect_faces(file_data, file_name)
        voice_result = await extract_voiceprint(file_data, file_name)
        status = 'done'
    else:
        status = 'done'

    import json as _json
    meta = _json.dumps({'file_type': ft, 'face': face_result, 'voice': voice_result, 'ingest': str(ingest_result)[:500]}, ensure_ascii=False, default=str)

    userorgid = await env.get_userorgid()
    from sqlor.dbpools import get_sor_context
    async with get_sor_context(env, 'rag') as sor:
        await sor.sqlExe(
            "INSERT INTO documents (id, kb_id, folder_id, file_name, file_type, file_size, file_path, mime_type, status, metadata, org_id, created_at, updated_at) "
            "VALUES (" + "${id}$, ${kb_id}$, ${folder_id}$, ${file_name}$, 'other', ${file_size}$, ${file_path}$, 'application/octet-stream', ${status}$, ${meta}$, ${org_id}$, NOW(), NOW())",
            {"id": doc_id, "kb_id": kb_id, "folder_id": folder_id, "file_name": file_name,
             "file_size": file_size, "file_path": "/idfile/files/" + saved_name,
             "status": status, "meta": meta, "org_id": userorgid})
        await sor.sqlExe(
            "UPDATE knowledge_bases SET doc_count=doc_count+1, total_size=total_size+" + "${size}$ WHERE id=${kb_id}$",
            {"size": file_size, "kb_id": kb_id})
        if ingest_result and status == 'done':
            chunks_n = ingest_result.get('chunks', 0) if isinstance(ingest_result, dict) else 0
            if chunks_n:
                await sor.sqlExe("UPDATE documents SET chunk_count=${chunks}$ WHERE id=${id}$", {"chunks": chunks_n, "id": doc_id})
                await sor.sqlExe("UPDATE knowledge_bases SET chunk_count=chunk_count+${n}$ WHERE id=${kb_id}$", {"n": chunks_n, "kb_id": kb_id})

    faces_n = face_result.get('faces', 0) if face_result and isinstance(face_result, dict) else 0
    speakers_n = voice_result.get('speakers', 0) if voice_result and isinstance(voice_result, dict) else 0
    return _json.dumps({
        "status": "SUCCEEDED", "doc_id": doc_id, "file_name": file_name,
        "file_size": file_size, "folder_id": folder_id, "file_type": ft,
        "rag_status": status, "faces": faces_n, "speakers": speakers_n
    }, ensure_ascii=False, default=str)
